"""ExamFlow RAG microservice adapted from he_thong_AI_sinh_de/main.py.
No legacy pickle index is loaded; each request gets its own in-memory FAISS index.
"""
import json
import asyncio
import logging
import math
import os
import re
from io import BytesIO
from typing import Literal

import httpx
import psycopg2
from dotenv import load_dotenv
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from langchain_community.vectorstores import FAISS
from langchain_ollama import OllamaEmbeddings
from langchain_core.documents import Document as LCDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel, Field
from pypdf import PdfReader

load_dotenv()
logger = logging.getLogger('examflow.duplicate')

OLLAMA_URL = os.getenv('OLLAMA_URL', 'http://127.0.0.1:11434').rstrip('/')
MODEL = os.getenv('OLLAMA_MODEL', 'qwen2.5:7b')
EMBED_MODEL = os.getenv('OLLAMA_EMBED_MODEL', 'nomic-embed-text')
MAX_CHARS = int(os.getenv('AI_MAX_SOURCE_CHARS', '120000'))
DUPLICATE_THRESHOLD = float(os.getenv('AI_DUPLICATE_THRESHOLD', '0.85'))
QUESTIONBANK_DSN = os.getenv('EXAMFLOW_QUESTIONBANK_DSN', '')
DUPLICATE_MAX_ATTEMPTS = int(os.getenv('AI_DUPLICATE_MAX_ATTEMPTS', '3'))
if not 1 <= DUPLICATE_MAX_ATTEMPTS <= 5:
    raise ValueError('AI_DUPLICATE_MAX_ATTEMPTS must be between 1 and 5')
if not 0.0 <= DUPLICATE_THRESHOLD <= 1.0:
    raise ValueError('AI_DUPLICATE_THRESHOLD must be between 0 and 1')
app = FastAPI(title='ExamFlow AI RAG')
app.add_middleware(CORSMiddleware, allow_origins=os.getenv('AI_CORS_ORIGINS', 'http://localhost:5173').split(','), allow_methods=['GET','POST'], allow_headers=['*'])

class Unit(BaseModel):
    name: str
    chapterId: int | None = None
    cloId: int | None = None
    nb: int = Field(ge=0, le=30)
    th: int = Field(ge=0, le=30)
    vd: int = Field(ge=0, le=30)

class GenerationRequest(BaseModel):
    courseId: int
    sourceText: str
    instruction: str = ''
    units: list[Unit]


def extract(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit('.', 1)[-1]
    if suffix == 'txt':
        return content.decode('utf-8-sig', errors='replace')
    if suffix == 'pdf':
        return '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(content)).pages)
    if suffix == 'docx':
        doc = Document(BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            paragraphs.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
        return '\n'.join(paragraphs)
    raise HTTPException(400, f'Không hỗ trợ file {filename}; dùng PDF/DOCX/TXT')

@app.get('/health')
def health():
    return {'status': 'ok', 'ollama_url': OLLAMA_URL, 'model': MODEL}

@app.post('/extract')
async def extract_documents(files: list[UploadFile] = File(...)):
    if not 1 <= len(files) <= 5:
        raise HTTPException(400, 'Chọn từ 1 đến 5 tài liệu')
    docs = []
    for file in files:
        raw = await file.read()
        if len(raw) > 20 * 1024 * 1024:
            raise HTTPException(413, f'{file.filename}: vượt quá 20 MB')
        try:
            content = extract(file.filename or '', raw).strip()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, f'Không đọc được {file.filename}: {exc}') from exc
        if not content:
            raise HTTPException(422, f'{file.filename}: không có văn bản; PDF scan cần OCR')
        docs.append({'name': file.filename, 'extractedText': content, 'wordCount': len(content.split())})
    return {'documents': docs}


def clean_json(s: str):
    s = re.sub(r'^\s*```(?:json)?\s*|\s*```\s*$', '', s.strip(), flags=re.I)
    return json.loads(s)


def normalize(response: str, bloom: str, unit: Unit):
    root = clean_json(response)
    questions = root.get('questions') if isinstance(root, dict) else root
    if not isinstance(questions, list):
        return []
    valid = []
    for q in questions:
        if not isinstance(q, dict):
            continue
        options = q.get('options')
        if isinstance(options, dict):
            options = [options.get(k) for k in 'ABCD']
        answer = q.get('correctIndex', q.get('answer'))
        if isinstance(answer, str) and answer.upper() in 'ABCD' and len(answer) == 1:
            answer = 'ABCD'.index(answer.upper())
        content = q.get('content') or q.get('question')
        if not isinstance(content, str) or not content.strip() or not isinstance(options, list) or len(options) != 4 or not all(isinstance(o, str) and o.strip() for o in options) or type(answer) is not int or answer not in range(4):
            continue
        valid.append({'content': content.strip(), 'options': options, 'correctIndex': answer, 'explanation': str(q.get('explanation') or ''), 'bloom': bloom, 'chapterId': unit.chapterId, 'cloId': unit.cloId})
    return valid

async def ollama_generate(client: httpx.AsyncClient, prompt: str) -> str:
    response = await client.post(f'{OLLAMA_URL}/api/generate', json={'model': MODEL, 'prompt': prompt, 'stream': False, 'format': 'json', 'options': {'num_ctx': 2048, 'num_predict': 850}})
    response.raise_for_status()
    return response.json().get('response', '')

def _nearest_questions(course_id: int, vector: list[float], limit: int = 3) -> list[dict]:
    """Read-only vector search. A question is compared only within its course."""
    if not QUESTIONBANK_DSN:
        raise RuntimeError('EXAMFLOW_QUESTIONBANK_DSN is not configured')
    if len(vector) != 768 or not all(math.isfinite(x) for x in vector):
        raise ValueError('Embedding must contain 768 finite values')
    vector_text = '[' + ','.join(str(x) for x in vector) + ']'
    with psycopg2.connect(QUESTIONBANK_DSN, connect_timeout=8) as conn:
        with conn.cursor() as cursor:
            cursor.execute('''
                SELECT q.id, q.content, 1 - (e.embedding <=> %s::extensions.vector) AS similarity
                FROM public.question_embeddings e
                JOIN public.questions q ON q.id = e.question_id
                WHERE q.course_id = %s AND e.model_name = %s
                ORDER BY e.embedding <=> %s::extensions.vector
                LIMIT %s
            ''', (vector_text, course_id, EMBED_MODEL, vector_text, limit))
            return [{'questionId': row[0], 'content': row[1], 'similarity': float(row[2])}
                    for row in cursor.fetchall()]


def _cosine(a: list[float], b: list[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    aa = math.sqrt(sum(x*x for x in a))
    bb = math.sqrt(sum(y*y for y in b))
    return dot / (aa * bb) if aa and bb else 0.0


async def _embed_question(text: str) -> list[float]:
    async with httpx.AsyncClient(timeout=90) as client:
        response = await client.post(f'{OLLAMA_URL}/api/embed', json={
            'model': EMBED_MODEL, 'input': text
        })
        response.raise_for_status()
        values = response.json()['embeddings'][0]
        if len(values) != 768:
            raise ValueError(f'Expected 768 embedding dimensions, got {len(values)}')
        return values


async def _duplicate_report(course_id: int, content: str, batch_vectors: list[list[float]]) -> tuple[dict, list[float]]:
    vector = await _embed_question(content)
    matches = await asyncio.to_thread(_nearest_questions, course_id, vector)
    batch_score = max((_cosine(vector, other) for other in batch_vectors), default=0.0)
    bank_score = matches[0]['similarity'] if matches else 0.0
    return ({'suspectedDuplicate': max(bank_score, batch_score) >= DUPLICATE_THRESHOLD,
             'threshold': DUPLICATE_THRESHOLD, 'bankMatches': matches,
             'maxBatchSimilarity': batch_score}, vector)


class DuplicateCheckRequest(BaseModel):
    courseId: int
    content: str = Field(min_length=3)


@app.post('/duplicate-check')
async def duplicate_check(request: DuplicateCheckRequest):
    """Explicit diagnostics endpoint; does not write any question or embedding."""
    try:
        report, _ = await _duplicate_report(request.courseId, request.content, [])
        return report
    except Exception as exc:
        logger.exception('Duplicate check unavailable')
        raise HTTPException(503, 'Duplicate checking unavailable; inspect ai-service logs') from exc


@app.post('/generate')
async def generate(request: GenerationRequest):
    total = sum(u.nb + u.th + u.vd for u in request.units)
    if total < 1 or total > 30:
        raise HTTPException(400, 'Số câu phải từ 1 đến 30')
    if len(request.sourceText.strip()) < 50:
        raise HTTPException(400, 'Nội dung tài liệu quá ngắn; vui lòng tải tài liệu thực')
    chunks = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=150).split_documents([LCDocument(page_content=request.sourceText[:MAX_CHARS])])
    try:
        embeddings = OllamaEmbeddings(base_url=OLLAMA_URL, model=EMBED_MODEL)
        store = await __import__('asyncio').to_thread(FAISS.from_documents, chunks, embeddings)
    except Exception as exc:
        raise HTTPException(502, f'Embedding/FAISS thất bại: {str(exc)[:250]}') from exc
    # Fail closed: never silently bypass duplicate detection during generation.
    if not QUESTIONBANK_DSN:
        raise HTTPException(503, 'Chưa cấu hình Question Bank để kiểm tra trùng')
    result = []
    seen = set()
    batch_vectors = []
    async with httpx.AsyncClient(timeout=httpx.Timeout(240.0, connect=10.0)) as client:
        for unit in request.units:
            for bloom, count in [('remember', unit.nb), ('understand', unit.th), ('apply', unit.vd)]:
                for index in range(count):
                    query = f'{unit.name} {request.instruction} {bloom}'
                    retrieved = await asyncio.to_thread(store.similarity_search, query, 4)
                    context = '\n---\n'.join(d.page_content for d in retrieved)[:4800]
                    prompt = f'''Bạn là giảng viên. Dựa CHỈ trên tài liệu dưới đây, tạo ĐÚNG 1 câu hỏi trắc nghiệm tiếng Việt.
Chủ đề: {unit.name}. Bloom: {bloom} (remember=nhớ, understand=hiểu, apply=vận dụng).
Yêu cầu bổ sung: {request.instruction}. Câu hỏi thứ {index+1} trong nhóm này.
Trả về JSON THUẦN đúng dạng {{"questions":[{{"content":"...","options":["...","...","...","..."],"correctIndex":0,"explanation":"..."}}]}}.
Đúng 4 đáp án, đúng một đáp án, không viết A/B/C/D trong options, giải thích có căn cứ trong tài liệu.
TÀI LIỆU:\n{context}'''
                    accepted = None
                    rejected_contents = []
                    last_error = ''
                    for attempt in range(DUPLICATE_MAX_ATTEMPTS):
                        retry_instruction = ''
                        if rejected_contents:
                            examples = '\n'.join(f'- {x[:250]}' for x in rejected_contents[-3:])
                            retry_instruction = (f'\nCÁC CÂU ĐÃ BỊ TỪ CHỐI VÌ TRÙNG/QUÁ GIỐNG:\n{examples}'
                                                 '\nHãy kiểm tra một kiến thức KHÁC trong tài liệu, không chỉ đổi cách diễn đạt.'
                                                 '\nChỉ xuất JSON hợp lệ, đúng một câu.')
                        try:
                            raw = await ollama_generate(client, prompt + retry_instruction)
                            candidates = normalize(raw, bloom, unit)
                        except (ValueError, httpx.HTTPError) as exc:
                            last_error = str(exc)
                            logger.warning('GENERATION_RETRY course=%s bloom=%s attempt=%s reason=model_or_json',
                                           request.courseId, bloom, attempt + 1)
                            continue
                        if not candidates:
                            last_error = 'AI trả JSON không có câu hỏi hợp lệ'
                            continue
                        for item in candidates:
                            fingerprint = re.sub(r'\W+', '', item['content'].lower())
                            if fingerprint in seen:
                                rejected_contents.append(item['content'])
                                last_error = 'Câu hỏi trùng nguyên văn trong cùng lượt'
                                continue
                            try:
                                report, vector = await _duplicate_report(request.courseId, item['content'], batch_vectors)
                            except Exception as exc:
                                logger.exception('DUPLICATE_CHECK_FAILED course=%s', request.courseId)
                                raise HTTPException(503, 'Không thể kiểm tra trùng; chưa trả câu hỏi để tránh bỏ qua kiểm tra') from exc
                            logger.info('DUPLICATE_AUDIT course=%s bloom=%s attempt=%s suspected=%s bank_top=%s batch_max=%.4f',
                                        request.courseId, bloom, attempt + 1, report['suspectedDuplicate'],
                                        [(m['questionId'], round(m['similarity'], 4)) for m in report['bankMatches']],
                                        report['maxBatchSimilarity'])
                            if report['suspectedDuplicate']:
                                rejected_contents.append(item['content'])
                                last_error = 'Câu hỏi nghi trùng ngữ nghĩa'
                                continue
                            accepted = item
                            seen.add(fingerprint)
                            batch_vectors.append(vector)  # Only accepted questions count in this batch.
                            break
                        if accepted is not None:
                            break
                    if accepted is None:
                        # No partial result: caller should not save an incomplete Bloom matrix.
                        logger.warning('GENERATION_EXHAUSTED course=%s bloom=%s index=%s attempts=%s reason=%s',
                                       request.courseId, bloom, index + 1, DUPLICATE_MAX_ATTEMPTS, last_error)
                        raise HTTPException(422, f'Không tạo được câu khác biệt cho mức {bloom} sau {DUPLICATE_MAX_ATTEMPTS} lần thử. Hãy bổ sung tài liệu hoặc giảm số câu.')
                    result.append(accepted)
    return result
